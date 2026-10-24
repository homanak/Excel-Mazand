# -*- coding: utf-8 -*-
"""Updates 00_Master.xlsx for the v03 series.

Adds the new dropdown sources that 02_Bank needs — bank operation types and the
"other weighbridge" persons — plus the dashboard entry on Home and a Help block.
"""
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from bank2 import BANK_OPS

FOREST = "FF14453D"
CREAM = "FFFBF8F1"
SAND2 = "FFF7F1E6"
GOLD = "FFC89B3C"
LINE = "FFDDD5C4"
WHITE = "FFFFFFFF"
MUTED = "FF6B7A76"
MINT = "FFE4F0E8"
TERRA = "FFC4623F"
FA = "B Nazanin"
EN = "Tahoma"

OPS = ["خرید مرکبات", "فروش بار (بارفروش)", "فروش متفرقه / صادراتی", "آبگیری",
       "انعام مصطفی", "باسکول دیگران", "درآمد پوز باسکول (بدون طرف)", "مرجوعی خرید",
       "مرجوعی فروش", "ضایعات / افت انبار", "انتقال بین انبار", "سایر ورود",
       "سایر خروج", "هزینه باسکول نقدی (درآمد)", "خرید سبد", "فروش سبد"]


def _f(size=10, bold=False, color="FF2B2A26", fill=None, border=False, align="right",
       wrap=True, name=FA):
    side = Side(style="thin", color=LINE) if border else None
    return dict(font=Font(name=name, size=size, bold=bold, color=color),
                fill=PatternFill("solid", fgColor=fill) if fill else None,
                border=Border(left=side, right=side, top=side, bottom=side) if border else None,
                alignment=Alignment(horizontal=align, vertical="center", wrap_text=wrap,
                                    readingOrder=2))


def style(cell, **kw):
    st = _f(**kw)
    cell.font = st["font"]
    if st["fill"]:
        cell.fill = st["fill"]
    if st["border"]:
        cell.border = st["border"]
    cell.alignment = st["alignment"]


def build(src, dst):
    wb = openpyxl.load_workbook(src)
    ls = wb["List"]

    # ---- read BD persons from Accounts
    acc = wb["Acc"]
    bd = []
    for r in range(10, 260):
        code = acc.cell(r, 1).value
        if code and str(code).startswith("BD"):
            bd.append((str(code), acc.cell(r, 4).value))
    banks = []
    bnk = wb["Bnk"]
    for r in range(10, 20):
        if bnk.cell(r, 1).value:
            banks.append((bnk.cell(r, 1).value, bnk.cell(r, 2).value))

    # ---- new list columns (S, T, U)
    start_col = 19   # S
    blocks = [
        ("نوع عملیات بانکی", "Bank operations", [o[0] for o in BANK_OPS]),
        ("شخص باسکول دیگران", "Other scales",
         [f"{c} – {n}" if n else f"{c} – (بدون نام)" for c, n in bd]),
        ("بانک عامل (نام)", "Company banks", [f"{c} – {t}" for c, t in banks]),
    ]
    for i, (fa, en, values) in enumerate(blocks):
        c = start_col + i
        letter = get_column_letter(c)
        head = ls.cell(8, c, fa)
        style(head, size=9.5, bold=True, color=WHITE, fill=FOREST, border=True, align="center")
        head2 = ls.cell(9, c, en)
        style(head2, size=7, color=MUTED, fill=SAND2, border=True, align="center")
        ls.column_dimensions[letter].width = 30 if i < 2 else 34
        for j, v in enumerate(values):
            cell = ls.cell(10 + j, c, v)
            style(cell, size=9.5, fill=CREAM, border=True, align="right")

    note_r = 32
    ls.merge_cells(start_row=note_r, start_column=start_col, end_row=note_r, end_column=start_col + 2)
    c = ls.cell(note_r, start_col,
                "این سه فهرست، کشوهای فایل 02_Bank را می‌سازند. برای اضافه کردن نوع عملیات بانکی جدید، "
                "یک ردیف به ستون S اضافه کنید و در شیت Sync فایل بانک همان ردیف را ادامه دهید.")
    style(c, size=9, color=MUTED, fill=SAND2, border=True)

    # ---- defined names
    want = {
        "ListBankOps": f"List!$S$10:$S$25",
        "ListOtherScalesLabel": f"List!$T$10:$T$29",
        "ListBankNames": f"List!$U$10:$U$19",
    }
    for name, ref in want.items():
        if name in wb.defined_names:
            del wb.defined_names[name]
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    # ---- Sync block (read-only mirror)
    sy = wb["Sync"]
    if sy.cell(6, 8).value in (None, ""):
        hdr = sy.cell(6, 8, "نوع عملیات بانکی")
        style(hdr, size=9.5, bold=True, color=WHITE, fill=FOREST, border=True, align="center")
        for j, (title, direction, mode, fixed, _n) in enumerate(BANK_OPS):
            sy.cell(7 + j, 8, title)
            style(sy.cell(7 + j, 8), size=9, fill=CREAM, border=True, align="right")
            sy.cell(7 + j, 9, direction)
            style(sy.cell(7 + j, 9), size=8.5, fill=WHITE, border=True, align="center")
            sy.cell(7 + j, 10, mode)
            style(sy.cell(7 + j, 10), size=8.5, fill=WHITE, border=True, align="center")
            sy.cell(7 + j, 11, fixed)
            style(sy.cell(7 + j, 11), size=8.5, fill=WHITE, border=True, align="center")
        sy.column_dimensions["H"].width = 26
        sy.column_dimensions["I"].width = 8
        sy.column_dimensions["J"].width = 9
        sy.column_dimensions["K"].width = 9

    # ---- Home: add dashboard card text
    home = wb["Home"]
    for r in range(1, 45):
        v = home.cell(r, 1).value
        if isinstance(v, str) and "فایل‌های دیگر مجموعه" in v:
            home.cell(r, 2).value = ("فایل‌های دیگر مجموعه  —  همه در همین پوشه  ·  داده‌ها از 01_Entry خوانده می‌شود  ·  "
                                     "داشبورد 08_Dash از ثبت و بانک هر دو می‌خواند")
            break

    # ---- Help block
    hp = wb["Help"]
    r = hp.max_row + 2
    hp.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
    c = hp.cell(r, 1, "۰۹   سری جدید (v03) — چه چیزهایی اضافه شد؟")
    style(c, size=12, bold=True, color=WHITE, fill=FOREST, border=True)
    lines = [
        "• فایل تراکنش‌ها (02_Bank) بازنویسی شد: کشوی بانک عامل، کد طرف حساب تایپی، انتقال داخلی با بانک مبدأ و مقصد، دفتر چک و پنج شیت تحلیل با نمودار.",
        "• داشبورد (08_Dash) اضافه شد: وزن، مبلغ، بانک‌ها، پز و طرف حساب‌ها در یک صفحه.",
        "• همه فهرست‌های کشویی در همین فایل تعریف شده‌اند؛ شیت List ستون‌های S، T و U سه فهرست جدید است.",
        "• همه شیت‌ها راست‌به‌چپ و آماده چاپ A4 هستند.",
    ]
    for i, txt in enumerate(lines):
        cell = hp.cell(r + 1 + i, 1, txt)
        hp.merge_cells(start_row=r + 1 + i, start_column=1, end_row=r + 1 + i, end_column=10)
        style(cell, size=10, fill=CREAM if i % 2 == 0 else WHITE, border=True)

    wb.calculation.fullCalcOnLoad = True
    wb.save(dst)
    return dst


def read_ops(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Op"]
    ops = []
    for row in ws.iter_rows(min_row=10, max_row=33, min_col=2, max_col=2, values_only=True):
        if row[0]:
            ops.append(row[0])
    wb.close()
    return ops
