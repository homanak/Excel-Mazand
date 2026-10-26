# -*- coding: utf-8 -*-
"""v05 — the report layer that the user asked for on top of the journals.

* ``LBRep``  (03_Report) — بارفروشان: one row per buyer with the totals, then
  every load of every buyer stacked underneath, dates one under another,
  ready to print.
* ``GrwTot`` (03_Report) — باغداران: volume total, ledger value, payments and
  an *invoice amount* column that can be typed by hand (or left alone when the
  purchase invoice already posted the value into the journal) plus a control
  that says whether the two agree.
* ``bill_final`` (01_Entry) — names the last totals line of the invoice for what
  it is: «مبلغ نهایی فاکتور».

Nothing here writes into the journal; every number is read from ``Data``
(03_Report) or from ``Book`` through the existing external links.
"""
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import look as L

FA, EN = L.FA, L.EN
DATA_LAST = 50009


# ----------------------------------------------------------------- helpers
def _dv(ws, ref, name, title, text):
    dv = DataValidation(type="list", formula1="=" + name, allow_blank=True,
                        showDropDown=False)
    dv.promptTitle = title
    dv.prompt = text
    dv.showInputMessage = True
    dv.errorTitle = "کد نامعتبر"
    dv.error = "کد را از فهرست انتخاب کنید."
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(ref)


def _band(ws, row, last, text, tint=L.FOREST, h=26):
    for c in range(1, last + 1):
        cell = ws.cell(row, c)
        cell.fill = PatternFill("solid", fgColor=tint)
        cell.border = L.box(tint)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
    c = ws.cell(row, 1)
    c.value = text
    c.font = Font(name=FA, size=11.5, bold=True, color=L.WHITE)
    c.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[row].height = h


def _head(ws, row, cols, tint=L.FOREST, h=32):
    for i, (t, w) in enumerate(cols, start=1):
        c = ws.cell(row, i, t)
        ws.column_dimensions[c.column_letter].width = w
        c.font = Font(name=FA, size=9, bold=True, color=L.WHITE)
        c.fill = PatternFill("solid", fgColor=tint)
        c.border = L.box(tint)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = h


def _cell(ws, row, col, value, fmt=None, name=None, bold=False, color=L.INK,
          bg=L.CREAM, h="center", size=9):
    c = ws.cell(row, col, value)
    c.font = Font(name=name or (EN if fmt else FA), size=size, bold=bold, color=color)
    c.fill = PatternFill("solid", fgColor=bg)
    c.border = L.box(L.LINE_S)
    c.alignment = Alignment(horizontal=h, vertical="center")
    if fmt:
        c.number_format = fmt
    return c


def _codes(master_path, group, cap):
    """The codes of one group, straight from the mother file (Acc)."""
    wb = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
    ws = wb["Acc"]
    out = []
    for row in ws.iter_rows(min_row=10, max_row=400, values_only=True):
        if row and row[1] == group and row[0]:
            out.append(str(row[0]).strip())
    wb.close()
    return (out + [""] * cap)[:cap]


# ------------------------------------------------------- 03_Report: LBRep
def lb_rep(wb, master_path):
    if "LBRep" in wb.sheetnames:
        del wb["LBRep"]
    idx = wb.sheetnames.index("BSel") + 1 if "BSel" in wb.sheetnames else 3
    ws = wb.create_sheet("LBRep", idx)
    last = 10
    L.brand(ws, last, "گزارش بارفروشان — جمع هر بارفروش و ریز بارها",
            "بالا: یک ردیف برای هر بارفروش. پایین: بارِ هر بارفروش با تاریخ‌ها زیر هم. "
            "آماده چاپ A4.", icon="◈")
    codes = _codes(master_path, "بارفروش", 20)

    # ---- ۰۱ summary
    _band(ws, 8, last, "۰۱   جمع هر بارفروش — یک ردیف برای هر نفر", tint=L.FOREST)
    _head(ws, 9, [("کد", 10), ("نام بارفروش", 26), ("زیرگروه", 15), ("تعداد بار", 9),
                  ("وزن کل (kg)", 13), ("مبلغ کل بار", 16), ("بدهکار", 15),
                  ("بستانکار", 15), ("مانده", 16), ("وضعیت", 22)])
    first = 10
    for n, code in enumerate(codes):
        r = first + n
        _cell(ws, r, 1, code or None, name=EN, bold=True, color=L.FOREST, bg=L.INPUT)
        _cell(ws, r, 2, f'=IF($A{r}="","",IFERROR(VLOOKUP($A{r},AcctTbl,2,0),"کد ناشناخته"))',
              h="right", bg=L.SAND)
        _cell(ws, r, 3, f'=IF($A{r}="","",IFERROR(VLOOKUP($A{r},AcctTbl,4,0),""))', h="right")
        _cell(ws, r, 4, f'=IF($A{r}="","",COUNTIF(Data!$C$10:$C${DATA_LAST},$A{r}))', fmt=L.INT)
        _cell(ws, r, 5, f'=IF($A{r}="","",SUMIF(Data!$C$10:$C${DATA_LAST},$A{r},'
                        f'Data!$M$10:$M${DATA_LAST}))', fmt=L.DEC)
        _cell(ws, r, 6, f'=IF($A{r}="","",SUMIF(Data!$C$10:$C${DATA_LAST},$A{r},'
                        f'Data!$I$10:$I${DATA_LAST}))', fmt=L.INT)
        _cell(ws, r, 7, f'=IF($A{r}="","",SUMIF(Data!$C$10:$C${DATA_LAST},$A{r},'
                        f'Data!$G$10:$G${DATA_LAST}))', fmt=L.INT, color=L.TERRA)
        _cell(ws, r, 8, f'=IF($A{r}="","",SUMIF(Data!$C$10:$C${DATA_LAST},$A{r},'
                        f'Data!$H$10:$H${DATA_LAST}))', fmt=L.INT, color=L.SAGE)
        _cell(ws, r, 9, f'=IF($A{r}="","",$G{r}-$H{r})', fmt=L.INT_RED, bold=True,
              color=L.FOREST, bg=L.MINT)
        _cell(ws, r, 10, f'=IF($A{r}="","",IF(ROUND($I{r},0)>0,"طلب ما از ایشان",'
                         f'IF(ROUND($I{r},0)<0,"بستانکاری ایشان","تسویه")))', h="right",
              bold=True, bg=L.MINT)
        ws.row_dimensions[r].height = 18
    tot = first + len(codes)
    ws.merge_cells(start_row=tot, start_column=1, end_row=tot, end_column=3)
    _cell(ws, tot, 1, "جمع بارفروشان", bold=True, color=L.WHITE, bg=L.FOREST, h="center")
    for col in (4, 5, 6, 7, 8):
        lt = "ABCDEFGHIJ"[col - 1]
        _cell(ws, tot, col, f"=SUM({lt}{first}:{lt}{tot - 1})", fmt=L.INT,
              bold=True, color=L.WHITE, bg=L.FOREST)
    for col in (9, 10):
        ws.cell(tot, col).fill = PatternFill("solid", fgColor=L.FOREST)
        ws.cell(tot, col).border = L.box(L.FOREST)
    _cell(ws, tot, 9, f"=$G{tot}-$H{tot}", fmt=L.INT_RED, bold=True, color=L.WHITE, bg=L.FOREST)
    _cell(ws, tot, 10, "", bg=L.FOREST)
    _dv(ws, f"A{first}:A{tot - 1}", "ListAllCodes", "کد بارفروش",
        "کد را از فهرست حساب‌ها انتخاب کنید؛ جای هر ردیف خالی هم می‌توانید کد بزنید.")

    # ---- ۰۲ detail, stacked by buyer and date
    start = tot + 2
    _band(ws, start, last, "۰۲   ریز بارها — بار هر بارفروش، تاریخ‌ها زیر هم", tint=L.INK)
    _head(ws, start + 1, [("کد", 10), ("نام بارفروش", 26), ("تاریخ", 12), ("سند", 12),
                          ("محصول", 14), ("وزن خالص (kg)", 13), ("مبلغ", 16),
                          ("بدهکار", 15), ("بستانکار", 15), ("مانده", 16)], tint=L.INK)
    d0 = start + 2
    N = 120
    key = f"SMALL(Data!$R$10:$R${DATA_LAST},ROW()-{d0 - 1})"
    pos = f"MATCH({key},Data!$R$10:$R${DATA_LAST},0)"
    for k in range(N):
        r = d0 + k
        bg = L.CREAM if k % 2 else L.WHITE
        _cell(ws, r, 1, f'=IFERROR(INDEX(Data!$C$10:$C${DATA_LAST},{pos}),"")',
              name=EN, bold=True, color=L.FOREST, bg=bg)
        _cell(ws, r, 2, f'=IF($A{r}="","",IF($A{r}=$A{r - 1},"",'
                        f'IFERROR(VLOOKUP($A{r},AcctTbl,2,0),"کد ناشناخته")))', h="right", bg=bg)
        _cell(ws, r, 3, f'=IF($A{r}="","",INDEX(Data!$A$10:$A${DATA_LAST},{pos}))', bg=bg)
        _cell(ws, r, 4, f'=IF($A{r}="","",INDEX(Data!$B$10:$B${DATA_LAST},{pos}))', bg=bg)
        _cell(ws, r, 5, f'=IF($A{r}="","",INDEX(Data!$K$10:$K${DATA_LAST},{pos}))', bg=bg)
        _cell(ws, r, 6, f'=IF($A{r}="","",INDEX(Data!$M$10:$M${DATA_LAST},{pos}))',
              fmt=L.DEC, bg=bg)
        _cell(ws, r, 7, f'=IF($A{r}="","",INDEX(Data!$I$10:$I${DATA_LAST},{pos}))',
              fmt=L.INT, bg=bg)
        _cell(ws, r, 8, f'=IF($A{r}="","",INDEX(Data!$G$10:$G${DATA_LAST},{pos}))',
              fmt=L.INT, color=L.TERRA, bg=bg)
        _cell(ws, r, 9, f'=IF($A{r}="","",INDEX(Data!$H$10:$H${DATA_LAST},{pos}))',
              fmt=L.INT, color=L.SAGE, bg=bg)
        _cell(ws, r, 10, f'=IF($A{r}="","",IF($A{r}<>$A{r - 1},0,N($J{r - 1}))+$H{r}-$I{r})',
              fmt=L.INT_RED, bold=True, color=L.FOREST, bg=L.MINT)
        ws.row_dimensions[r].height = 16
    dt = d0 + N
    ws.merge_cells(start_row=dt, start_column=1, end_row=dt, end_column=5)
    _cell(ws, dt, 1, "جمع همین صفحه", bold=True, color=L.WHITE, bg=L.FOREST)
    for col in (6, 7, 8, 9):
        lt = "ABCDEFGHIJ"[col - 1]
        _cell(ws, dt, col, f"=SUM({lt}{d0}:{lt}{dt - 1})", fmt=L.INT, bold=True,
              color=L.WHITE, bg=L.FOREST)
    _cell(ws, dt, 10, "", bg=L.FOREST)

    for col, w in zip("ABCDEFGHIJ", (10, 26, 13, 12, 14, 13, 16, 15, 15, 17)):
        ws.column_dimensions[col].width = w

    L.note(ws, dt + 2, last,
           "ردیف‌های بالا از دفتر «ثبت روزانه» خوانده می‌شوند: هر بارِ خارج‌شده از این بارفروش "
           "یک ردیف است، به ترتیب بارفروش و بعد تاریخ — تاریخ‌ها زیر هم. نام هر بارفروش فقط "
           "در ردیف اول تکرار می‌شود. اگر جدول خالی بود، در 03_Report شیت Data را با "
           "Edit Links → Update Values تازه کنید.", bg=L.MINT, h=40)
    L.page(ws, last, dt + 2, landscape=True)
    return "LBRep"


# ------------------------------------------------------ 03_Report: GrwTot
def grw_tot(wb, master_path):
    if "GrwTot" in wb.sheetnames:
        del wb["GrwTot"]
    idx = wb.sheetnames.index("LBRep") + 1 if "LBRep" in wb.sheetnames else 4
    ws = wb.create_sheet("GrwTot", idx)
    last = 11
    L.brand(ws, last, "باغداران — وزن، مبلغ خرید و صورت‌حساب",
            "بار ورودی سورتینگ پول ندارد؛ ارزش خرید در ستون «صورت‌حساب» نوشته می‌شود — "
            "دستی یا از جمع فاکتور خرید. ستون کنترل می‌گوید با دفتر می‌خواند یا نه.",
            icon="◈", tint=L.INK)
    L.kpis(ws, 6, last, [
        ("جمع وزن (kg)", f"=SUM($E$10:$E$69)", L.DEC, L.SAGE),
        ("مبلغ دفتر", f"=SUM($F$10:$F$69)", L.INT, L.TERRA),
        ("پرداختی‌ها", f"=SUM($G$10:$G$69)", L.INT, L.GOLD),
        ("جمع صورت‌حساب", f"=SUM($I$10:$I$69)", L.INT, L.FOREST),
    ])
    _band(ws, 8, last, "۰۱   یک ردیف برای هر باغدار — وزن، پرداخت، صورت‌حساب و مانده",
          tint=L.FOREST, h=24)
    _head(ws, 9, [("کد", 9), ("نام باغدار", 26), ("زیرگروه", 13), ("تعداد بار", 8),
                  ("وزن کل (kg)", 12), ("مبلغ خرید در دفتر", 16), ("پرداختی به باغدار", 16),
                  ("مانده دفتر", 15), ("مبلغ صورت‌حساب (دستی)", 16),
                  ("مانده با صورت‌حساب", 16), ("کنترل", 20)])
    codes = _codes(master_path, "باغدار", 60)
    first = 10
    for n, code in enumerate(codes):
        r = first + n
        bg = L.CREAM if n % 2 else L.WHITE
        _cell(ws, r, 1, code or None, name=EN, bold=True, color=L.FOREST, bg=L.INPUT)
        _cell(ws, r, 2, f'=IF($A{r}="","",IFERROR(VLOOKUP($A{r},AcctTbl,2,0),"کد ناشناخته"))',
              h="right", bg=L.SAND)
        _cell(ws, r, 3, f'=IF($A{r}="","",IFERROR(VLOOKUP($A{r},AcctTbl,4,0),""))', h="right", bg=bg)
        _cell(ws, r, 4, f'=IF($A{r}="","",COUNTIF(Data!$C$10:$C${DATA_LAST},$A{r}))', fmt=L.INT, bg=bg)
        _cell(ws, r, 5, f'=IF($A{r}="","",SUMIF(Data!$C$10:$C${DATA_LAST},$A{r},'
                        f'Data!$M$10:$M${DATA_LAST}))', fmt=L.DEC, bg=bg)
        _cell(ws, r, 6, f'=IF($A{r}="","",SUMIF(Data!$C$10:$C${DATA_LAST},$A{r},'
                        f'Data!$H$10:$H${DATA_LAST}))', fmt=L.INT, color=L.TERRA, bg=bg)
        _cell(ws, r, 7, f'=IF($A{r}="","",SUMIF(Data!$C$10:$C${DATA_LAST},$A{r},'
                        f'Data!$G$10:$G${DATA_LAST}))', fmt=L.INT, color=L.GOLD, bg=bg)
        _cell(ws, r, 8, f'=IF($A{r}="","",$F{r}-$G{r})', fmt=L.INT_RED, bold=True,
              color=L.FOREST, bg=L.MINT)
        _cell(ws, r, 9, None, fmt=L.INT, bg=L.INPUT)
        _cell(ws, r, 10, f'=IF(OR($A{r}="",$I{r}=""),"",$I{r}-$G{r})', fmt=L.INT_RED,
              bold=True, color=L.FOREST, bg=L.MINT)
        _cell(ws, r, 11, f'=IF($A{r}="","",IF($I{r}="","— بدون صورت‌حساب دستی",'
                         f'IF(ABS($I{r}-$F{r})<=1,"✓ مطابق دفتر","⚠ اختلاف با دفتر")))',
              h="right", bold=True, bg=bg)
        ws.row_dimensions[r].height = 17
    tot = first + len(codes)
    ws.merge_cells(start_row=tot, start_column=1, end_row=tot, end_column=3)
    _cell(ws, tot, 1, "جمع باغداران", bold=True, color=L.WHITE, bg=L.FOREST)
    for col in (4, 5, 6, 7, 8, 9, 10):
        lt = "ABCDEFGHIJK"[col - 1]
        _cell(ws, tot, col, f"=SUM({lt}{first}:{lt}{tot - 1})", fmt=L.INT, bold=True,
              color=L.WHITE, bg=L.FOREST)
    _cell(ws, tot, 11, "", bg=L.FOREST)
    _dv(ws, f"A{first}:A{tot - 1}", "ListGrowerCodes", "کد باغدار",
        "کد باغدار را از فهرست انتخاب کنید.")

    L.note(ws, tot + 2, last,
           "بار ورودی سورتینگ فقط وزن دارد و در دفتر بدون مبلغ ثبت می‌شود؛ همین است که "
           "ستون «مبلغ خرید در دفتر» برای باغدارهای سورتینگی صفر می‌ماند. مبلغ واقعی خرید را "
           "در ستون زرد «مبلغ صورت‌حساب» بنویسید (یا جمع ردیف نهایی فاکتور خرید همان باغدار "
           "را از شیت BillBuy بردارید). ستون کنترل نشان می‌دهد صورت‌حساب با دفتر می‌خواند یا نه.",
           bg=L.MINT, h=44)
    L.page(ws, last, tot + 2, landscape=True)
    return "GrwTot"


# ------------------------------------------------- 01_Entry: invoice last row
def bill_final(path):
    """Rename the last totals line so the invoice clearly ends with an amount."""
    wb = openpyxl.load_workbook(path)
    out = 0
    for name, old, new in (
            ("Bill", "خالص قابل فروش پس از ضایعات", "مبلغ نهایی فاکتور — خالص پس از ضایعات"),
            ("BillBuy", "خالص قابل فروش پس از ضایعات", "مبلغ نهایی فاکتور — خالص پس از ضایعات")):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for row in ws.iter_rows(min_row=20, max_row=40, max_col=3):
            for c in row:
                if isinstance(c.value, str) and c.value.strip() == old:
                    c.value = new
                    out += 1
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    wb.close()
    return out


# --------------------------------------------------------------- the report
def report(path, master_path):
    global DATA_LAST
    wb = openpyxl.load_workbook(path)
    if "Data" in wb.sheetnames:
        ws = wb["Data"]
        last = max(ws.max_row, 10)
        DATA_LAST = last
        for r in range(10, last + 1):
            ws.cell(r, 18).value = (
                f'=IF($E{r}<>"بارفروش","",'
                f'IFERROR(MATCH($C{r},ListLoadBuyerCodes,0),0)*100000000'
                f'+IFERROR(VALUE(SUBSTITUTE(LEFT($A{r},10),"/","")),0)+ROW()/100000)')
        ws.column_dimensions["R"].hidden = True
    lb_rep(wb, master_path)
    grw_tot(wb, master_path)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    wb.close()
    return ["LBRep", "GrwTot"]
