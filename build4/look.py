# -*- coding: utf-8 -*-
"""openpyxl design layer — the same visual language as ``build2/kit.py`` but
usable on workbooks that are edited in place (01_Entry, 03…06, 07)."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FOREST = "FF14453D"
FOREST_2 = "FF1E5D52"
FOREST_3 = "FF0D312C"
SAGE = "FF7FA99B"
SAGE_L = "FFD8E6DF"
MINT = "FFE4F0E8"
CREAM = "FFFBF8F1"
SAND = "FFF3EADB"
SAND2 = "FFF7F1E6"
GOLD = "FFC89B3C"
GOLD_L = "FFEFDFB8"
TERRA = "FFC4623F"
BLUSH = "FFF2DCD3"
SKY = "FFDCE9F2"
SKY_D = "FF2B5C7A"
LAV = "FFE7E2F2"
GREY = "FFEEECE6"
LINE = "FFDDD5C4"
LINE_S = "FFEDE6D8"
INK = "FF2B2A26"
MUTED = "FF6B7A76"
WHITE = "FFFFFFFF"
RED = "FFB4544A"
GREEN = "FF3F7A57"
BLUE = "FF2B5C7A"
INPUT = "FFFFF4D6"

FA = "B Nazanin"
FA_T = "B Titr"
EN = "Tahoma"
INT = "#,##0"
INT_RED = "#,##0;[Red]-#,##0"
DEC = "#,##0.00"
PCT = "0.0%"

COMPANY = "مازندشت  ·  سورتینگ و بسته‌بندی مرکبات"
COMPANY_EN = "MAZANDASHT  ·  CITRUS SORTING & PACKING"


def side(color=LINE, style="thin"):
    return Side(style=style, color=color)


def fill(color):
    return PatternFill("solid", fgColor=color)


def font(name=FA, size=10, bold=False, color=INK, italic=False):
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)


def align(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, readingOrder=2)


def box(color=LINE):
    s = side(color)
    return Border(left=s, right=s, top=s, bottom=s)


def st(cell, name=FA, size=10, bold=False, color=INK, bg=None, border=LINE,
       h="center", v="center", wrap=False, fmt=None, italic=False):
    cell.font = font(name, size, bold, color, italic)
    if bg:
        cell.fill = fill(bg)
    if border:
        cell.border = box(border)
    cell.alignment = align(h, v, wrap)
    if fmt:
        cell.number_format = fmt
    return cell


def put(ws, ref, value, **kw):
    c = ws[ref]
    c.value = value
    st(c, **kw)
    return c


# ------------------------------------------------------------------ sections
def brand(ws, last, title, sub, icon="", tint=FOREST, rights=None):
    """Rows 1-5: dark ribbon, big title, subtitle, gold line."""
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 40
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 6
    ws.row_dimensions[5].height = 6
    for c in range(1, last + 1):
        for r in (1, 2, 3):          # the ribbon is dark: white type needs it
            ws.cell(r, c).fill = fill(tint)
        ws.cell(4, c).fill = fill(GOLD)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last)
    c1 = ws.cell(1, 1)
    c1.value = f"{COMPANY}    ·    {COMPANY_EN}"
    c1.font = font(EN, 8, True, SAGE_L)
    c1.alignment = align("right")
    ws.cell(2, 1).value = ("    " + icon + "  " if icon else "") + title
    ws.cell(2, 1).font = font(FA_T, 22, True, WHITE)
    ws.cell(2, 1).alignment = align("right")
    ws.cell(3, 1).value = sub
    ws.cell(3, 1).font = font(FA, 11, False, SAGE_L)
    ws.cell(3, 1).alignment = align("right")


def section(ws, row, last, num, title, sub="", tint=FOREST):
    for i in range(1, last + 1):
        cell = ws.cell(row, i)
        cell.fill = fill(tint)
        cell.border = Border(bottom=side(tint))
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(1, last - 1))
    c = ws.cell(row, 1)
    c.value = f"{title}   {sub}".strip()
    c.font = font(FA, 12, True, WHITE)
    c.alignment = align("right")
    if last > 1:
        chip = ws.cell(row, last)
        chip.value = num
        chip.font = font(EN, 9, True, GOLD_L)
        chip.alignment = align("center")
    ws.row_dimensions[row].height = 26


def thead(ws, row, cols, tint=FOREST, height=34):
    """cols = [(fa, en, width, tooltip|None)]"""
    for i, spec in enumerate(cols, start=1):
        fa, en, w = spec[0], spec[1], spec[2]
        tip = spec[3] if len(spec) > 3 else None
        ws.column_dimensions[ws.cell(row, i).column_letter].width = w
        c = ws.cell(row, i)
        c.value = fa
        c.font = font(FA, 9, True, WHITE)
        c.fill = fill(tint)
        c.border = box(tint)
        c.alignment = align("center", "center", True)
        if tip:
            from openpyxl.comments import Comment
            c.comment = Comment(tip, "MAZANDASHT", width=280, height=110)
    ws.row_dimensions[row].height = height


def note(ws, row, last, text, bg=SAND, h=34, color=INK, size=9.5):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
    for i in range(1, last + 1):
        ws.cell(row, i).fill = fill(bg)
    c = ws.cell(row, 1)
    c.value = "   " + text
    c.font = font(FA, size, False, color)
    c.alignment = align("right", "center", True)
    c.border = Border(bottom=side(LINE_S))
    ws.row_dimensions[row].height = h


def kpis(ws, row, last, items):
    """items = [(label, formula, fmt, color)] — equal spans across ``last``."""
    n = len(items)
    bounds, c = [], 1
    base, extra = divmod(last, n)
    for i in range(n):
        w = base + (1 if i < extra else 0)
        bounds.append((c, min(c + w - 1, last)))
        c += w
    for (label, formula, fmt, color), (c0, c1) in zip(items, bounds):
        if c1 > c0:
            ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c1)
        lab = ws.cell(row, c0)
        lab.value = label
        lab.font = font(FA, 8, True, MUTED)
        lab.fill = fill(SAND2)
        lab.alignment = align("center")
        for i in range(c0, c1 + 1):
            ws.cell(row, i).fill = fill(SAND2)
            ws.cell(row, i).border = box(LINE)
        if c1 > c0:
            ws.merge_cells(start_row=row + 1, start_column=c0, end_row=row + 1, end_column=c1)
        v = ws.cell(row + 1, c0)
        v.value = formula
        v.font = font(EN, 16, True, color)
        v.alignment = align("center")
        v.number_format = fmt
        for i in range(c0, c1 + 1):
            ws.cell(row + 1, i).fill = fill(CREAM)
            ws.cell(row + 1, i).border = box(LINE)
    ws.row_dimensions[row].height = 20
    ws.row_dimensions[row + 1].height = 32


def page(ws, last_col, last_row, landscape=True, header_rows=None, margins=0.3):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = margins
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.5
    ws.oddHeader.left.text = "مازندشت — سامانه حسابداری سورتینگ"
    ws.oddHeader.left.size = 8
    ws.oddHeader.center.text = "MAZANDASHT · " + ws.title.upper()
    ws.oddHeader.center.size = 8
    ws.oddHeader.right.text = "&A"
    ws.oddHeader.right.size = 8
    ws.oddFooter.left.text = "&F"
    ws.oddFooter.left.size = 8
    ws.oddFooter.center.text = "صفحه &P از &N"
    ws.oddFooter.center.size = 8
    ws.oddFooter.right.text = "نسخه ۰۴ · ۱۴۰۵"
    ws.oddFooter.right.size = 8
    from openpyxl.utils import get_column_letter
    ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    if header_rows:
        ws.print_title_rows = header_rows


def tip(ws, cell, text, w=280, h=110):
    from openpyxl.comments import Comment
    ws[cell].comment = Comment(text, "MAZANDASHT", width=w, height=h)
