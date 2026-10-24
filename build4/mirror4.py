# -*- coding: utf-8 -*-
"""Report files (03/04/05/06) — mirror the 50,000-line journal, keep A4 print
geometry intact and add a print header/footer to every sheet."""
import openpyxl
from openpyxl.utils import get_column_letter

import grow
import look as L


def polish(wb, name=""):
    """A4 + fit-to-width + header/footer for every sheet, without touching
    the print areas that are already set."""
    n = 0
    for ws in wb.worksheets:
        try:
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.print_options.horizontalCentered = True
            ws.oddHeader.left.text = "مازندشت — سامانه حسابداری سورتینگ"
            ws.oddHeader.center.text = f"MAZANDASHT · {ws.title.upper()}"
            ws.oddHeader.right.text = "&A"
            ws.oddFooter.left.text = "&F"
            ws.oddFooter.center.text = "صفحه &P از &N"
            ws.oddFooter.right.text = "نسخه ۰۴ · ۱۴۰۵"
            for part in (ws.oddHeader.left, ws.oddHeader.center, ws.oddHeader.right,
                         ws.oddFooter.left, ws.oddFooter.center, ws.oddFooter.right):
                part.size = 8
            if not ws.print_area:
                ws.print_area = (f"A1:{get_column_letter(ws.max_column)}"
                                 f"{min(ws.max_row, 50010)}")
            n += 1
        except Exception:
            continue
    return n


def build(path, mirror=True):
    wb = openpyxl.load_workbook(path)
    if "Data" in wb.sheetnames and mirror:
        ws = wb["Data"]
        grow.grow_data(ws, grow.DATA_LAST)
        grow.reseat_ranges(wb, {1807: grow.DATA_LAST, 1808: grow.DATA_LAST,
                                1809: grow.DATA_LAST})
    sheets = polish(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return sheets
