# -*- coding: utf-8 -*-
"""Renders one worksheet to a PNG preview (no Excel / LibreOffice needed).

    python build4/preview.py Mazandasht_v04/01_Entry.xlsx Bill /tmp/bill.png

It understands merged cells, fills, borders, fonts, alignment and the row/column
geometry, draws right-to-left when the sheet is RTL, and reshapes Persian text.
Good enough to check the layout of the printed invoices before opening Excel.
"""
import os
import sys

import openpyxl
from openpyxl.utils import get_column_letter

try:
    import arabic_reshaper
    from bidi.algorithm import get_display
except Exception:                                     # pragma: no cover
    arabic_reshaper = None

from PIL import Image, ImageDraw, ImageFont

SCALE = 1.6
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def _shape(text):
    if arabic_reshaper is None or not any("\u0600" <= c <= "\u06ff" for c in text):
        return text
    try:
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def _rgb(fill):
    if fill is None or fill.fill_type is None:
        return None
    c = getattr(fill.fgColor, "rgb", None)
    if isinstance(c, str) and len(c) == 8 and c != "00000000":
        return "#" + c[2:]
    return None


def _font(cell, bold_default=False):
    f = cell.font
    size = int(round((f.sz or 11) * 1.18 * SCALE))
    bold = bool(f.b) or bold_default
    path = FONT_B if bold else FONT
    try:
        return ImageFont.truetype(path, max(7, size))
    except Exception:
        return ImageFont.load_default()


def _text(cell):
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, str) and v.startswith("="):
        return "‹محاسبه خودکار›"
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return f"{v:,}"
    return str(v)


def render(path, sheet, out, cols=None, rows=None, pad=24, bg="#FFFFFF"):
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]

    if cols is None:
        cols = [get_column_letter(i) for i in range(1, ws.max_column + 1)]
    if rows is None:
        rows = range(1, ws.max_row + 1)

    widths = {}
    for c in cols:
        d = ws.column_dimensions.get(c)
        w = d.width if d is not None and d.width else 9
        widths[c] = max(20, w * 7.4 + 6) * SCALE
    heights = {}
    for r in rows:
        d = ws.row_dimensions.get(r)
        heights[r] = (d.height if d is not None and d.height else 15) * 1.34 * SCALE

    W = int(sum(widths.values())) + pad * 2
    H = int(sum(heights.values())) + pad * 2
    img = Image.new("RGB", (W, H), bg)
    d = ImageDraw.Draw(img)

    x0 = {}
    if ws.sheet_view.rightToLeft:
        right = pad + sum(widths.values())
        for c in cols:                     # column A sits on the right edge
            right -= widths[c]
            x0[c] = right
    else:
        x = pad
        for c in cols:
            x0[c] = x
            x += widths[c]

    y0 = {}
    y = pad
    for r in rows:
        y0[r] = y
        y += heights[r]

    merged = {}
    for mr in ws.merged_cells.ranges:
        letters = [get_column_letter(i) for i in range(mr.min_col, mr.max_col + 1)]
        letters = [c for c in letters if c in x0]
        rs = [r for r in range(mr.min_row, mr.max_row + 1) if r in y0]
        if not letters or not rs:
            continue
        left = min(x0[c] for c in letters)
        right = max(x0[c] + widths[c] for c in letters)
        top = y0[min(rs)]
        bottom = max(y0[r] + heights[r] for r in rs)
        merged[(mr.min_row, mr.min_col)] = (left, top, right, bottom)

    for r in rows:
        for c in cols:
            cell = ws[f"{c}{r}"]
            ci = openpyxl.utils.column_index_from_string(c)
            box = merged.get((r, ci))
            if box is None:
                left = x0[c]
                top = y0[r]
                right = left + widths[c]
                bottom = top + heights[r]
            else:
                left, top, right, bottom = box

            fill = _rgb(cell.fill)
            if fill:
                d.rectangle([left, top, right - 1, bottom - 1], fill=fill, outline=None)
            if cell.border is not None and cell.border.left is not None \
                    and cell.border.left.style:
                d.rectangle([left, top, right - 1, bottom - 1], outline="#E3DDD0")

            text = _text(cell)
            if not text:
                continue
            font = _font(cell)
            lines = text.split("\n")
            al = (cell.alignment.horizontal or "general")
            va = cell.alignment.vertical or "bottom"
            lh = font.size + 3 * SCALE
            total_h = lh * len(lines)
            ty = top + (bottom - top - total_h) / 2 if va == "center" else top + 3
            for line in lines:
                shown = _shape(line)
                w = d.textlength(shown, font=font)
                if al in ("center", "justify"):
                    tx = left + (right - left - w) / 2
                elif al == "right" or ws.sheet_view.rightToLeft and al == "general":
                    tx = right - w - 3
                elif al == "left":
                    tx = left + 3
                else:
                    tx = left + (right - left - w) / 2
                color = "#" + cell.font.color.rgb[2:] if (
                    cell.font.color is not None and isinstance(cell.font.color.rgb, str)
                    and len(cell.font.color.rgb) == 8) else "#222222"
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    color = "#8C8677"
                d.text((tx, ty), shown, font=font, fill=color)
                ty += lh

    d.rectangle([0, 0, W - 1, H - 1], outline="#BBB6AA")
    img.save(out)
    return out


if __name__ == "__main__":
    args = sys.argv[1:]
    wb_path = args[0]
    sheet = args[1]
    out = args[2] if len(args) > 2 else "/tmp/preview.png"
    limited = args[3] if len(args) > 3 else None
    cols = rows = None
    if limited:
        cpart, rpart = limited.split(":")
        cols = [get_column_letter(i) for i in range(1, openpyxl.utils.column_index_from_string(cpart) + 1)]
        rows = range(1, int(rpart) + 1)
    print(render(wb_path, sheet, out, cols, rows))
